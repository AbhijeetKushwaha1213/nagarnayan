"""
NAGAR NAYAN - SIH 2026 Idea Presentation Generator
Style: Reference PPT (Telhan Sathi) | Structure: SIH 2026 Template (6 slides)
Install:  pip install python-pptx
Run:      python nagar_nayan_ppt.py   ->  NagarNayan_SIH2026.pptx
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ---------------- COLORS ----------------
BLUE      = RGBColor(0x1F, 0x63, 0xB0)   # header blue
NAVY      = RGBColor(0x1B, 0x3A, 0x6B)   # title dark blue
DARK      = RGBColor(0x20, 0x20, 0x20)
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
LIGHTBLUE = RGBColor(0xE8, 0xF0, 0xFA)   # panel bg
YELLOW    = RGBColor(0xFF, 0xD9, 0x66)   # flow boxes
SKY       = RGBColor(0xBD, 0xE3, 0xF5)
GREEN     = RGBColor(0x2E, 0x7D, 0x32)
DGREEN    = RGBColor(0x4E, 0x6E, 0x3D)
RED       = RGBColor(0xD6, 0xB8, 0x28) if False else RGBColor(0xD6, 0x28, 0x28)
ORANGE    = RGBColor(0xF2, 0x7B, 0x21)
GRAYBG    = RGBColor(0xF2, 0xF2, 0xF2)
STEEL     = RGBColor(0x4A, 0x5A, 0x6A)

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

# ---------------- HELPERS ----------------
def add_slide():
    return prs.slides.add_slide(BLANK)

def rect(s, x, y, w, h, fill, line=None, shape=MSO_SHAPE.RECTANGLE, line_w=1.0, shadow=False):
    sp = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line; sp.line.width = Pt(line_w)
    sp.shadow.inherit = shadow
    return sp

def txt(s, x, y, w, h, paras, anchor=MSO_ANCHOR.TOP, wrap=True):
    """paras: list of (align, space_after_pt, [(text, size, bold, color, italic), ...])"""
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = wrap; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Pt(2); tf.margin_top = tf.margin_bottom = Pt(1)
    for i, (align, sa, runs) in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.space_after = Pt(sa)
        for t, size, bold, color, italic in runs:
            r = p.add_run(); r.text = t
            r.font.size = Pt(size); r.font.bold = bold
            r.font.color.rgb = color; r.font.italic = italic
            r.font.name = "Calibri"
    return tb

def P(text, size=14, bold=False, color=DARK, italic=False, align=PP_ALIGN.LEFT, sa=4):
    return (align, sa, [(text, size, bold, color, italic)])

def PR(runs, align=PP_ALIGN.LEFT, sa=4):   # rich paragraph
    return (align, sa, runs)

def set_text_in_shape(sp, text, size, bold, color, align=PP_ALIGN.CENTER):
    tf = sp.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = Pt(3); tf.margin_top = tf.margin_bottom = Pt(1)
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color
    r.font.name = "Calibri"

def flow_box(s, x, y, w, h, text, fill=YELLOW, tcolor=DARK, size=12, bold=True, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    sp = rect(s, x, y, w, h, fill, line=DARK, shape=shape, line_w=1.0)
    set_text_in_shape(sp, text, size, bold, tcolor)
    return sp

def down_arrow(s, cx, y, h=0.28):
    rect(s, cx - 0.11, y, 0.22, h, DARK, shape=MSO_SHAPE.DOWN_ARROW)

def num_circle(s, x, y, n, color):
    c = rect(s, x, y, 0.42, 0.42, color, shape=MSO_SHAPE.OVAL)
    set_text_in_shape(c, n, 13, True, WHITE)

def sih_logo(s, x, y, scale=1.0):
    """Recreates SIH 2026 logo (hexagon + brain) — replace with official PNG if available."""
    u = scale
    rect(s, x, y, 1.55*u, 1.45*u, RGBColor(0xE9,0xE9,0xE9), shape=MSO_SHAPE.HEXAGON)
    rect(s, x+0.22*u, y+0.28*u, 0.50*u, 0.62*u, ORANGE, shape=MSO_SHAPE.OVAL)
    g = rect(s, x+0.74*u, y+0.28*u, 0.50*u, 0.62*u, GREEN, shape=MSO_SHAPE.OVAL)
    set_text_in_shape(g, "01", 8, True, WHITE)
    t = rect(s, x+0.42*u, y+0.98*u, 0.70*u, 0.32*u, WHITE)
    set_text_in_shape(t, "SIH", 12, True, STEEL)
    txt(s, x+1.65*u, y+0.10*u, 2.0*u, 0.9*u,
        [PR([("SMART INDIA\n", 12*scale, True, STEEL, False)]),
         PR([("HACKATHON\n", 12*scale, True, STEEL, False)]),
         PR([("2026", 12*scale, True, STEEL, False)])], wrap=False)

def footer(s, num):
    rect(s, 0, 7.08, 13.333, 0.42, BLUE)
    txt(s, 0, 7.10, 13.333, 0.38,
        [P("SMART INDIA HACKATHON 2026", 15, True, WHITE, align=PP_ALIGN.CENTER)], anchor=MSO_ANCHOR.MIDDLE)
    txt(s, 12.5, 7.10, 0.7, 0.38, [P(str(num), 14, True, WHITE, align=PP_ALIGN.RIGHT)], anchor=MSO_ANCHOR.MIDDLE)

def team_oval(s):
    o = rect(s, 0.25, 0.15, 1.55, 0.85, WHITE, line=BLUE, shape=MSO_SHAPE.OVAL, line_w=1.5)
    set_text_in_shape(o, "Your Team\nName", 12, False, STEEL)

def slide_title(s, text):
    txt(s, 2.0, 0.10, 9.5, 0.85, [P(text, 36, True, DARK, align=PP_ALIGN.CENTER,
        )], anchor=MSO_ANCHOR.MIDDLE)
    s.shapes[-1].text_frame.paragraphs[0].runs[0].font.name = "Times New Roman"

# ============================================================
# SLIDE 1 — TITLE PAGE
# ============================================================
s = add_slide()
txt(s, 0.4, 0.25, 10.2, 1.0, [P("SMART INDIA HACKATHON 2026", 44, True, NAVY)])
s.shapes[-1].text_frame.paragraphs[0].runs[0].font.name = "Times New Roman"
sih_logo(s, 10.9, 0.25)

# big hexagon visual (right side, like reference)
rect(s, 8.6, 1.6, 4.4, 4.6, RGBColor(0xEC,0xEC,0xEC), shape=MSO_SHAPE.HEXAGON)
rect(s, 9.35, 2.35, 1.30, 1.75, ORANGE, shape=MSO_SHAPE.OVAL)
g = rect(s, 10.55, 2.35, 1.30, 1.75, GREEN, shape=MSO_SHAPE.OVAL)
set_text_in_shape(g, "0 1 0 1\n1 0 1 0", 11, True, WHITE)
t = rect(s, 9.85, 4.25, 1.55, 0.55, WHITE)
set_text_in_shape(t, "SIH", 22, True, STEEL)

txt(s, 3.4, 1.35, 6.6, 0.8, [P("TITLE PAGE", 30, True, DARK, align=PP_ALIGN.CENTER)])
s.shapes[-1].text_frame.paragraphs[0].runs[0].font.name = "Times New Roman"

items = [
    ("Problem Statement ID – ", ""),
    ("Problem Statement Title – ", ""),
    ("Theme – ", ""),
    ("PS Category – Software", ""),
    ("Team ID – ", ""),
    ("Team Name – ", ""),
]
paras = [PR([(label, 20, True, DARK, False), (val, 20, False, DARK, False)], sa=14)
         for label, val in items]
txt(s, 0.5, 2.45, 7.6, 4.0, paras)

# tagline ribbon
r = rect(s, 1.2, 6.45, 6.2, 0.5, GRAYBG, line=STEEL, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
set_text_in_shape(r, "Every Bus Becomes an Eye on the City.", 15, True, BLUE)

# ============================================================
# SLIDE 2 — IDEA TITLE & PROPOSED SOLUTION
# ============================================================
s = add_slide()
team_oval(s); sih_logo(s, 10.95, 0.18, scale=0.85)
txt(s, 2.0, 0.05, 9.0, 0.65, [P("NAGAR NAYAN", 34, True, DARK, align=PP_ALIGN.CENTER)])
s.shapes[-1].text_frame.paragraphs[0].runs[0].font.name = "Times New Roman"
txt(s, 2.0, 0.62, 9.0, 0.45, [P("Eyes of the City — AI-Powered Mobile Urban Intelligence",
    17, False, DARK, italic=True, align=PP_ALIGN.CENTER)])

# ---- Left panel : Proposed Solution ----
hb = rect(s, 0.55, 1.20, 5.6, 0.55, BLUE, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
set_text_in_shape(hb, "Proposed Solution / Approach", 20, True, WHITE)
txt(s, 0.18, 1.18, 0.4, 0.5, [P("💡", 22, False, DARK)])

panel = rect(s, 0.55, 1.90, 5.6, 3.55, LIGHTBLUE, line=BLUE, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
txt(s, 0.75, 2.00, 5.2, 3.4, [
    PR([("The Problem: ", 14, True, NAVY, False),
        ("Cities depend on fixed CCTV cameras with limited coverage, manual road inspections "
         "and citizen complaints — causing incomplete city-wide visibility and slow response.", 14, False, DARK, False)], sa=8),
    PR([("Our Solution: ", 14, True, NAVY, False),
        ("Nagar Nayan transforms existing public bus cameras into a distributed AI urban "
         "sensing network that continuously detects road, traffic, infrastructure and safety events.", 14, False, DARK, False)], sa=8),
    PR([("Detected Events:", 14, True, NAVY, False)], sa=4),
])

# detection chips (2 rows x 3)
chips = ["🕳 Road Defects", "🌊 Waterlogging", "🚦 Infrastructure",
         "🚗 Traffic Jams", "🚶 Pedestrian Risk", "⚠ Safety Incidents"]
for i, c in enumerate(chips):
    cx = 0.75 + (i % 3) * 1.78; cy = 4.15 + (i // 3) * 0.55
    ch = rect(s, cx, cy, 1.68, 0.45, YELLOW, line=DARK, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    set_text_in_shape(ch, c, 10.5, True, DARK)

# ---- Right panel : workflow diagram ----
wf = rect(s, 6.45, 1.20, 6.55, 4.25, WHITE, line=DARK, shape=MSO_SHAPE.ROUNDED_RECTANGLE, line_w=1.5)
txt(s, 6.45, 1.22, 6.55, 0.4, [P("NAGAR NAYAN End-to-End Workflow", 16, True, DARK, align=PP_ALIGN.CENTER)])
cx = 8.55  # center of left column of flow
flow_box(s, cx-1.3, 1.70, 2.6, 0.42, "🚌 BUS CAMERAS + GPS")
down_arrow(s, cx, 2.12)
flow_box(s, cx-1.3, 2.42, 2.6, 0.42, "📡 LIVE STREAM SERVER (RTSP)")
down_arrow(s, cx, 2.84)
flow_box(s, cx-1.3, 3.14, 2.6, 0.48, "🤖 AI VISION ENGINE", fill=BLUE, tcolor=WHITE, size=12)
down_arrow(s, cx, 3.62)
flow_box(s, cx-1.95, 3.92, 1.55, 0.42, "Road Analysis", fill=SKY, size=10.5)
flow_box(s, cx+0.40, 3.92, 1.55, 0.42, "Traffic Analysis", fill=SKY, size=10.5)
down_arrow(s, cx, 4.36)
flow_box(s, cx-1.3, 4.66, 2.6, 0.42, "⚡ Event Validation → GIS Command Center", size=11)

# side mini-steps
txt(s, 11.05, 1.70, 1.9, 3.4, [
    P("🎥 Frame\nSampling", 11, True, NAVY, align=PP_ALIGN.CENTER, sa=10),
    P("✅ Multi-frame\nValidation", 11, True, NAVY, align=PP_ALIGN.CENTER, sa=10),
    P("☁ Backend\nFastAPI", 11, True, NAVY, align=PP_ALIGN.CENTER, sa=10),
    P("🗺 Map Dashboard", 11, True, NAVY, align=PP_ALIGN.CENTER),
])

# ---- Innovation & Uniqueness banner ----
ib = rect(s, 5.1, 5.05, 3.0, 0.5, DGREEN, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
set_text_in_shape(ib, "Innovation and Uniqueness", 15, True, WHITE)
card = rect(s, 0.55, 5.60, 12.25, 1.25, RGBColor(0xFB,0xE9,0xD9), line=ORANGE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, line_w=2)
txt(s, 0.75, 5.68, 11.9, 1.1, [
    PR([("USP: ", 14, True, NAVY, False),
        ("Instead of installing thousands of new sensors and cameras, Nagar Nayan intelligently reuses the "
         "existing moving camera infrastructure of public buses — zero new hardware, dynamic coverage of every "
         "road a bus travels, at city scale.", 14, False, DARK, False)], sa=2),
    PR([("♻ Low cost  •  📡 Always moving  •  🏙 City-wide dynamic coverage", 13, True, ORANGE, False)],
       align=PP_ALIGN.CENTER, sa=0),
])
footer(s, 2)

# ============================================================
# SLIDE 3 — TECHNICAL APPROACH
# ============================================================
s = add_slide()
team_oval(s); sih_logo(s, 10.95, 0.18, scale=0.85)
txt(s, 2.0, 0.08, 9.0, 0.7, [P("TECHNICAL APPROACH", 36, True, DARK, align=PP_ALIGN.CENTER)])
s.shapes[-1].text_frame.paragraphs[0].runs[0].font.name = "Times New Roman"

zones = [
    ("Zone 1\nData Capture", ["🚌 Bus Camera(s)", "📍 GPS Module", "🎞 Recorded / Live Feed"], RGBColor(0xE9,0xF5,0xFB)),
    ("Zone 2\nStreaming Layer", ["📡 MediaMTX Server", "🎞 FFmpeg Pipeline", "🔗 RTSP / WebRTC"], RGBColor(0xE9,0xF8,0xEE)),
    ("Zone 3\nAI Vision Engine", ["🎯 YOLO Detection", "👁 OpenCV + Tracking", "✅ Multi-frame Validation", "🛣 Road + 🚦 Traffic"], RGBColor(0xFB,0xF3,0xE4)),
    ("Zone 4\nCommand Center", ["⚡ FastAPI + WebSockets", "🗄 PostgreSQL + PostGIS", "🗺 React + MapLibre GIS", "📊 Live Event Dashboard"], RGBColor(0xF1,0xEC,0xFB)),
]
zx = 0.35
for zname, zitems, zcolor in zones:
    rect(s, zx, 1.05, 2.42, 4.55, zcolor, line=STEEL)
    rect(s, zx, 1.05, 2.42, 0.62, RGBColor(0xDD,0xE7,0xF0), line=STEEL)
    txt(s, zx, 1.08, 2.42, 0.6, [P(zname, 13.5, True, NAVY, align=PP_ALIGN.CENTER)], anchor=MSO_ANCHOR.MIDDLE)
    yy = 1.85
    for it in zitems:
        b = rect(s, zx+0.15, yy, 2.12, 0.72, YELLOW, line=DARK, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        set_text_in_shape(b, it, 11, True, DARK)
        yy += 0.92
    if zx < 7.0:
        rect(s, zx+2.44, 3.05, 0.26, 0.30, DARK, shape=MSO_SHAPE.RIGHT_ARROW)
    zx += 2.62

# ---- Implementation Process (right) ----
ip = rect(s, 10.85, 1.05, 2.35, 4.55, WHITE, line=DARK, shape=MSO_SHAPE.ROUNDED_RECTANGLE, line_w=1.5)
txt(s, 10.90, 1.08, 2.25, 0.45, [P("⚙ Implementation Process", 13, True, RED, align=PP_ALIGN.CENTER)])
steps = [("Stream Capture", "Bus video + GPS ingested live"),
         ("Frame Sampling", "Only key frames processed"),
         ("AI Detection", "YOLO flags road/traffic events"),
         ("Validation", "Multi-frame + confidence filter"),
         ("Event Generation", "Geo-tagged, deduplicated events"),
         ("GIS Visualization", "Live map at Command Center")]
yy = 1.55
for i, (t1, t2) in enumerate(steps, 1):
    num_circle(s, 10.95, yy, str(i), ORANGE if i % 2 else BLUE)
    txt(s, 11.42, yy-0.06, 1.75, 0.7, [
        PR([(t1, 11, True, DARK, False)], sa=0),
        PR([(t2, 9, False, STEEL, False)], sa=0)])
    yy += 0.68

# ---- Tech stack bar ----
ts = rect(s, 0.35, 5.75, 12.65, 1.10, WHITE, line=DARK, shape=MSO_SHAPE.ROUNDED_RECTANGLE, line_w=1.5)
txt(s, 0.45, 5.80, 2.1, 1.0, [
    PR([("Components /\n", 14, True, BLUE, False)], sa=0),
    PR([("Technology\n", 14, True, BLUE, False)], sa=0),
    PR([("stack to be used", 14, True, BLUE, False)], sa=0)])
stack = [("Python", RGBColor(0x37,0x76,0xAB)), ("OpenCV", RGBColor(0x5C,0x3D,0x2E)),
         ("YOLO", RGBColor(0xC6,0x28,0x28)), ("MediaMTX", RGBColor(0x2E,0x7D,0x32)),
         ("FastAPI", RGBColor(0x00,0x96,0x88)), ("PostGIS", RGBColor(0x33,0x66,0x99)),
         ("React", RGBColor(0x21,0x96,0xF3)), ("MapLibre", RGBColor(0x51,0x4A,0x9E)),
         ("WebSockets", RGBColor(0x60,0x7D,0x8B))]
xx = 2.75
for name, col in stack:
    c = rect(s, xx, 6.02, 1.12, 0.52, col, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    set_text_in_shape(c, name, 11, True, WHITE)
    xx += 1.14
txt(s, 0.45, 6.55, 12.5, 0.3,
    [P("Processing Strategy:  Live Stream → Frame Sampling → AI Detection → Multi-frame Validation → Event Generation → GIS Visualization",
       11.5, True, NAVY, align=PP_ALIGN.CENTER)])
footer(s, 3)

# ============================================================
# SLIDE 4 — FEASIBILITY AND VIABILITY
# ============================================================
s = add_slide()
team_oval(s); sih_logo(s, 10.95, 0.18, scale=0.85)
txt(s, 2.0, 0.08, 9.0, 0.7, [P("FEASIBILITY AND VIABILITY", 36, True, DARK, align=PP_ALIGN.CENTER)])
s.shapes[-1].text_frame.paragraphs[0].runs[0].font.name = "Times New Roman"

tops = [
    ("Feasibility", [
        "Existing bus cameras — no new sensing hardware needed.",
        "Modular architecture: Streaming → AI → Backend → Dashboard scale independently.",
        "Low-cost prototype: recorded videos, simulated RTSP streams, open-source AI models."]),
    ("Viability", [
        "Runs on existing public transport fleet — proven, always-on infrastructure.",
        "Software-only build keeps cost low and deployment fast.",
        "Geo-tagged event data directly useful for municipal maintenance planning."]),
    ("Practical Implementation", [
        "Prototype ready with single bus stream, then multi-bus scale-up.",
        "Containerized services enable city-scale rollout.",
        "Works with intermittent connectivity via stream retry & recovery."]),
]
xx = 0.35
for title, pts in tops:
    rect(s, xx, 1.00, 4.18, 1.90, GRAYBG, line=STEEL)
    txt(s, xx+0.12, 1.04, 3.95, 0.4, [P(title, 17, True, BLUE)])
    paras = [PR([("• ", 11.5, True, DARK, False), (p, 11.5, False, DARK, False)], sa=3) for p in pts]
    txt(s, xx+0.12, 1.45, 3.95, 1.45, paras)
    xx += 4.24

# ---- Challenges vs Strategies ----
big = rect(s, 0.35, 3.05, 12.65, 3.35, RGBColor(0xFA,0xFA,0xF5), line=DARK, shape=MSO_SHAPE.ROUNDED_RECTANGLE, line_w=1.5)
# left circle
c = rect(s, 0.75, 4.30, 1.9, 1.35, GRAYBG, shape=MSO_SHAPE.OVAL, line=STEEL)
set_text_in_shape(c, "⚠ Potential Challenges & Risks", 12, True, DARK)
challenges = [
    ("High video processing load", "Frame sampling + optimized AI inference"),
    ("False detections", "Multi-frame validation & confidence scoring"),
    ("Duplicate reports", "GPS + spatial event clustering (PostGIS)"),
    ("Bandwidth consumption", "Process only required frames/events"),
    ("Poor network connectivity", "Stream retry & recovery mechanisms"),
]
strategies = [
    ("01", 3.35), ("02", 3.90), ("03", 4.45), ("04", 5.00), ("05", 5.55),
]
yy = 3.25
for i, (ch, st) in enumerate(challenges, 1):
    num_circle(s, 2.85, yy, f"{i:02d}", RED)
    txt(s, 3.32, yy-0.08, 3.55, 0.6, [
        PR([(ch+": ", 11.5, True, DARK, False), (st, 11.5, False, DARK, False)], sa=0)])
    yy += 0.545
sc = rect(s, 11.05, 4.30, 1.8, 1.35, GRAYBG, shape=MSO_SHAPE.OVAL, line=STEEL)
set_text_in_shape(sc, "✔ Strategies For Overcoming Challenges", 11.5, True, DARK)
pairs = ["Frame sampling & lightweight inference cuts compute load.",
         "Multi-frame validation eliminates false alarms.",
         "Spatial clustering merges duplicate events into one report.",
         "Only key frames and event metadata are transmitted.",
         "Auto-reconnect + buffering handles patchy networks."]
yy = 3.25
for i, st in enumerate(pairs, 1):
    num_circle(s, 10.50, yy, f"{i:02d}", GREEN)
    txt(s, 7.05, yy-0.08, 3.40, 0.6, [PR([(st, 11.5, False, DARK, False)], align=PP_ALIGN.RIGHT, sa=0)])
    yy += 0.545

# deployment strip
dep = rect(s, 0.35, 6.52, 12.65, 0.48, RGBColor(0xDE,0xEB,0xF7), line=BLUE, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
set_text_in_shape(dep, "Deployment:  Prototype  →  Single Stream  →  Multiple Cameras  →  Multiple Buses  →  City-Scale Rollout",
                  13, True, NAVY)
footer(s, 4)

# ============================================================
# SLIDE 5 — IMPACT AND BENEFITS
# ============================================================
s = add_slide()
team_oval(s); sih_logo(s, 10.95, 0.18, scale=0.85)
txt(s, 2.0, 0.08, 9.0, 0.7, [P("IMPACT AND BENEFITS", 36, True, DARK, align=PP_ALIGN.CENTER)])
s.shapes[-1].text_frame.paragraphs[0].runs[0].font.name = "Times New Roman"

# central donut
rect(s, 2.55, 2.55, 2.5, 2.5, ORANGE, shape=MSO_SHAPE.DONUT)
cc = rect(s, 2.80, 2.95, 2.0, 1.7, WHITE, shape=MSO_SHAPE.OVAL)
set_text_in_shape(cc, "Potential Impact on Targeted Audience", 12, True, NAVY)

def aud(x, y, w, h, title, body):
    rect(s, x, y, w, h, GRAYBG, line=STEEL)
    txt(s, x+0.1, y+0.03, w-0.2, h-0.06, [
        PR([(title+"\n", 13, True, BLUE, False)], sa=1),
        PR([(body, 11, False, DARK, False)], sa=0)])

aud(0.45, 1.30, 3.9, 1.15, "🏙 City Authorities",
    "Real-time visibility of urban problems, faster defect identification, data-driven maintenance planning and better traffic management.")
aud(3.95, 1.30, 3.7, 1.15, "👥 Citizens",
    "Safer roads, faster resolution of infrastructure problems, reduced accident risks and better public mobility.")
aud(0.45, 5.15, 3.9, 1.15, "💰 Economic",
    "Lower manual inspection costs, early detection prevents expensive road damage, optimized maintenance budgets.")
aud(3.95, 5.15, 3.7, 1.15, "🌱 Environmental",
    "Less inspection travel, smoother traffic flow reduces idle emissions, data-driven planning supports sustainable mobility.")

# right benefits panel
bb = rect(s, 8.0, 1.10, 5.0, 0.55, BLUE, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
set_text_in_shape(bb, "Benefits of the Solution", 18, True, WHITE)
benefits = [
    ("Social", "Safer streets and faster grievance redressal build citizen trust in city governance.", RGBColor(0xE8,0xF0,0xFA)),
    ("Economic", "No new hardware capex; reuses existing cameras and cuts recurring inspection costs.", RGBColor(0xFB,0xF3,0xE4)),
    ("Environmental", "Reduced vehicle idling and smarter maintenance routes lower the city's carbon footprint.", RGBColor(0xE9,0xF8,0xEE)),
]
yy = 1.80
for name, body, col in benefits:
    rect(s, 8.0, yy, 5.0, 1.30, col, line=ORANGE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, line_w=1.5)
    txt(s, 8.15, yy+0.05, 4.7, 1.2, [
        PR([(name+"\n", 15, True, NAVY, False)], sa=2),
        PR([('"' + body + '"', 11.5, False, DARK, True)], sa=0)])
    yy += 1.45

# reactive -> proactive strip
rect(s, 4.55, 2.85, 3.1, 0.55, RGBColor(0xDE,0xEB,0xF7), line=BLUE, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
txt(s, 4.55, 2.87, 3.1, 0.5, [
    PR([("REACTIVE CITY → ", 11, True, RED, False)], align=PP_ALIGN.CENTER, sa=0),
    PR([("PROACTIVE CITY", 11, True, GREEN, False)], align=PP_ALIGN.CENTER, sa=0)])

q = rect(s, 0.45, 6.45, 12.5, 0.55, GRAYBG, line=STEEL, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
set_text_in_shape(q, '"Nagar Nayan helps cities see problems before they become crises."', 15, True, NAVY)
footer(s, 5)

# ============================================================
# SLIDE 6 — RESEARCH AND REFERENCES
# ============================================================
s = add_slide()
team_oval(s); sih_logo(s, 10.95, 0.18, scale=0.85)
txt(s, 2.0, 0.08, 9.0, 0.7, [P("RESEARCH AND REFERENCES", 36, True, DARK, align=PP_ALIGN.CENTER)])
s.shapes[-1].text_frame.paragraphs[0].runs[0].font.name = "Times New Roman"

rect(s, 0.45, 1.05, 12.45, 5.75, WHITE, line=DARK, shape=MSO_SHAPE.ROUNDED_RECTANGLE, line_w=1.5)

txt(s, 0.70, 1.15, 11.9, 2.1, [
    PR([("Research Areas Explored", 18, True, BLUE, False)], sa=6),
    PR([("• ", 13, True, DARK, False),
        ("Intelligent Transportation Systems — ", 13, True, DARK, False),
        ("AI-powered traffic monitoring and urban mobility analysis.", 13, False, DARK, False)], sa=4),
    PR([("• ", 13, True, DARK, False),
        ("Automated Road Damage Detection — ", 13, True, DARK, False),
        ("Computer-vision approaches for detecting potholes and road-surface defects (RDD datasets & papers).", 13, False, DARK, False)], sa=4),
    PR([("• ", 13, True, DARK, False),
        ("Object Detection & Tracking — ", 13, True, DARK, False),
        ("Real-time detection of vehicles, pedestrians and infrastructure using deep learning.", 13, False, DARK, False)], sa=4),
    PR([("• ", 13, True, DARK, False),
        ("Geospatial Urban Intelligence — ", 13, True, DARK, False),
        ("GIS-based visualization and spatial clustering for city management.", 13, False, DARK, False)], sa=4),
])

txt(s, 0.70, 3.45, 11.9, 3.2, [
    PR([("Academic & Technical Sources", 18, True, BLUE, False)], sa=6),
    PR([("• ", 13, True, DARK, False),
        ("YOLO — Real-Time Object Detection: ", 13, True, DARK, False),
        ("https://docs.ultralytics.com", 13, False, RGBColor(0x05,0x63,0xC1), False)], sa=4),
    PR([("• ", 13, True, DARK, False),
        ("OpenCV — Video Processing: ", 13, True, DARK, False),
        ("https://opencv.org", 13, False, RGBColor(0x05,0x63,0xC1), False)], sa=4),
    PR([("• ", 13, True, DARK, False),
        ("MediaMTX — RTSP Media Streaming: ", 13, True, DARK, False),
        ("https://github.com/bluenviron/mediamtx", 13, False, RGBColor(0x05,0x63,0xC1), False)], sa=4),
    PR([("• ", 13, True, DARK, False),
        ("FFmpeg — Video Streaming & Processing: ", 13, True, DARK, False),
        ("https://ffmpeg.org", 13, False, RGBColor(0x05,0x63,0xC1), False)], sa=4),
    PR([("• ", 13, True, DARK, False),
        ("PostgreSQL + PostGIS — Geospatial Intelligence: ", 13, True, DARK, False),
        ("https://postgis.net", 13, False, RGBColor(0x05,0x63,0xC1), False)], sa=4),
    PR([("• ", 13, True, DARK, False),
        ("OpenStreetMap / MapLibre — Map Data & Rendering: ", 13, True, DARK, False),
        ("https://www.openstreetmap.org  |  https://maplibre.org", 13, False, RGBColor(0x05,0x63,0xC1), False)], sa=4),
    PR([("• ", 13, True, DARK, False),
        ("Smart City Surveillance & Sensing research and Road Damage Detection datasets (RDD2022, crowdsourced road-imagery benchmarks).",
         13, False, DARK, False)], sa=4),
])
footer(s, 6)

output_path = os.path.join(os.path.dirname(__file__), "NagarNayan_SIH2026.pptx")
prs.save(output_path)
print(f"✅ Saved: {output_path} (6 slides)")
